# -*- coding: utf-8 -*-
"""
exam_builder_base.py
Base utility to generate professional Vietnamese examination documents in DOCX format.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

NAVY_BLUE = RGBColor(26, 54, 93)     # #1A365D
DARK_SLATE = RGBColor(45, 55, 72)    # #2D3748
CRIMSON = RGBColor(155, 44, 44)      # #9B2C2C
BLACK = RGBColor(0, 0, 0)
GRAY_TEXT = RGBColor(100, 116, 139)

def create_styled_document():
    doc = docx.Document()
    
    # Configure margins (Top 2cm, Bottom 2cm, Left 2.5cm, Right 2.0cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.79)     # ~2.0 cm
        section.bottom_margin = Inches(0.79)  # ~2.0 cm
        section.left_margin = Inches(0.98)    # ~2.5 cm
        section.right_margin = Inches(0.79)   # ~2.0 cm
        
        # Add page numbering in footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("Trang ")
        r_ft.font.name = "Times New Roman"
        r_ft.font.size = Pt(10)
        r_ft.font.italic = True
        r_ft.font.color.rgb = GRAY_TEXT
        
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_ft._p.append(fldSimple)
        
        r_slash = p_ft.add_run(" / ")
        r_slash.font.name = "Times New Roman"
        r_slash.font.size = Pt(10)
        r_slash.font.italic = True
        r_slash.font.color.rgb = GRAY_TEXT
        
        fldSimpleTotal = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        p_ft._p.append(fldSimpleTotal)

    # Set normal style font
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = BLACK
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(3)

    return doc

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_padding(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0BEC5", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_exam_header(doc, school_name, exam_title, subject_title, duration_str, exam_code="101"):
    """
    Creates a standard two-column official Vietnamese examination header.
    """
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(tbl)
    
    # Left column: Institution info
    cell_left = tbl.cell(0, 0)
    cell_left.width = Inches(3.2)
    p_l1 = cell_left.paragraphs[0]
    p_l1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l1.paragraph_format.space_after = Pt(2)
    p_l1.paragraph_format.line_spacing = 1.15
    r = p_l1.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = False
    
    r2 = p_l1.add_run(school_name.upper() + "\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11.5)
    r2.font.bold = True
    r2.font.color.rgb = NAVY_BLUE
    
    r3 = p_l1.add_run("--------------------")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(9)
    r3.font.color.rgb = GRAY_TEXT

    # Right column: Exam info
    cell_right = tbl.cell(0, 1)
    cell_right.width = Inches(3.6)
    p_r1 = cell_right.paragraphs[0]
    p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r1.paragraph_format.space_after = Pt(2)
    p_r1.paragraph_format.line_spacing = 1.15
    
    r_ex = p_r1.add_run(exam_title.upper() + "\n")
    r_ex.font.name = "Times New Roman"
    r_ex.font.size = Pt(12)
    r_ex.font.bold = True
    r_ex.font.color.rgb = CRIMSON
    
    r_sub = p_r1.add_run(f"MÔN: {subject_title.upper()}\n")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = NAVY_BLUE
    
    r_time = p_r1.add_run(f"Thời gian làm bài: {duration_str}\n(Không kể thời gian phát đề)\n")
    r_time.font.name = "Times New Roman"
    r_time.font.size = Pt(10)
    r_time.font.italic = True
    
    r_code = p_r1.add_run(f"MÃ ĐỀ THI: {exam_code}")
    r_code.font.name = "Times New Roman"
    r_code.font.size = Pt(10.5)
    r_code.font.bold = True

    # Student metadata box & Score box
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(8)
    p_meta.paragraph_format.space_after = Pt(4)
    r_meta = p_meta.add_run("Họ và tên thí sinh: ............................................................................ Số báo danh: ............................. Lớp: ...................")
    r_meta.font.name = "Times New Roman"
    r_meta.font.size = Pt(11)
    r_meta.font.italic = True

    # Score table
    score_tbl = doc.add_table(rows=2, cols=3)
    score_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(score_tbl, color="718096", sz="6")
    
    headers = [("ĐIỂM SỐ", Inches(1.8)), ("HỌ TÊN, CHỮ KÝ GIÁM KHẢO", Inches(2.5)), ("LỜI PHÊ CỦA THẦY / CÔ GIÁO", Inches(2.7))]
    for i, (head_text, w) in enumerate(headers):
        c = score_tbl.cell(0, i)
        c.width = w
        set_cell_background(c, "F1F5F9")
        set_cell_padding(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(head_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.bold = True
        
        # Second row empty for written notes
        c2 = score_tbl.cell(1, i)
        c2.width = w
        set_cell_padding(c2, top=200, bottom=200, left=100, right=100)
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(24)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(6)

def add_section_header(doc, section_title, note_str=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(section_title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = NAVY_BLUE
    
    if note_str:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_after = Pt(4)
        r_n = p_note.add_run(note_str)
        r_n.font.name = "Times New Roman"
        r_n.font.size = Pt(11)
        r_n.font.italic = True
        r_n.font.color.rgb = DARK_SLATE

def add_mcq_question(doc, q_number, q_text, opt_a, opt_b, opt_c, opt_d):
    """
    Adds a 4-choice multiple choice question with standard indentation.
    """
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.left_indent = Inches(0.15)
    
    r_num = p_q.add_run(f"Câu {q_number}: ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(11.5)
    r_num.font.bold = True
    
    r_txt = p_q.add_run(q_text)
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(11.5)
    
    # Options layout
    max_len = max(len(opt_a), len(opt_b), len(opt_c), len(opt_d))
    if max_len < 32:
        # All on one or two lines
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.space_after = Pt(3)
        p_opt.paragraph_format.left_indent = Inches(0.35)
        
        runs = [("A. ", opt_a, "    \t"), ("B. ", opt_b, "    \t"),
                ("C. ", opt_c, "    \t"), ("D. ", opt_d, "")]
        for prefix, txt, tab in runs:
            r_p = p_opt.add_run(prefix)
            r_p.font.name = "Times New Roman"
            r_p.font.size = Pt(11.5)
            r_p.font.bold = True
            r_t = p_opt.add_run(txt + tab)
            r_t.font.name = "Times New Roman"
            r_t.font.size = Pt(11.5)
    elif max_len < 55:
        # 2 rows of 2 options
        for pair in [(("A. ", opt_a), ("B. ", opt_b)), (("C. ", opt_c), ("D. ", opt_d))]:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_after = Pt(2)
            p_opt.paragraph_format.left_indent = Inches(0.35)
            for prefix, txt in pair:
                r_p = p_opt.add_run(prefix)
                r_p.font.name = "Times New Roman"
                r_p.font.size = Pt(11.5)
                r_p.font.bold = True
                r_t = p_opt.add_run(txt.ljust(45) + "  ")
                r_t.font.name = "Times New Roman"
                r_t.font.size = Pt(11.5)
    else:
        # 4 separate lines
        for prefix, txt in [("A. ", opt_a), ("B. ", opt_b), ("C. ", opt_c), ("D. ", opt_d)]:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_after = Pt(2)
            p_opt.paragraph_format.left_indent = Inches(0.35)
            r_p = p_opt.add_run(prefix)
            r_p.font.name = "Times New Roman"
            r_p.font.size = Pt(11.5)
            r_p.font.bold = True
            r_t = p_opt.add_run(txt)
            r_t.font.name = "Times New Roman"
            r_t.font.size = Pt(11.5)

def add_true_false_question(doc, q_number, context_passage, statements):
    """
    Format new 2025 BGD format: Reading context + 4 statements (a, b, c, d)
    """
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(6)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.left_indent = Inches(0.15)
    
    r_num = p_q.add_run(f"Câu {q_number}: ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(11.5)
    r_num.font.bold = True
    
    r_inst = p_q.add_run("Đọc đoạn tư liệu / bảng thông tin sau đây và chọn Đúng (Đ) hoặc Sai (S) cho từng ý a), b), c), d):\n")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(11.5)
    
    # Context quote box
    p_box = doc.add_paragraph()
    p_box.paragraph_format.left_indent = Inches(0.35)
    p_box.paragraph_format.right_indent = Inches(0.35)
    p_box.paragraph_format.space_after = Pt(4)
    r_ctx = p_box.add_run(f'"{context_passage}"')
    r_ctx.font.name = "Times New Roman"
    r_ctx.font.size = Pt(11)
    r_ctx.font.italic = True
    r_ctx.font.color.rgb = DARK_SLATE
    
    # Table for 4 statements with checkbox columns
    tbl = doc.add_table(rows=len(statements)+1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="4")
    
    headers = [("Lệnh hỏi / Mệnh đề", Inches(5.2)), ("Đúng", Inches(0.9)), ("Sai", Inches(0.9))]
    for col_idx, (head, w) in enumerate(headers):
        c = tbl.cell(0, col_idx)
        c.width = w
        set_cell_background(c, "F8FAFC")
        set_cell_padding(c, top=60, bottom=60, left=80, right=80)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(head)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)
        r.font.bold = True
        
    for row_idx, (letter, stmt_txt) in enumerate(statements, start=1):
        c0 = tbl.cell(row_idx, 0)
        c0.width = Inches(5.2)
        set_cell_padding(c0, top=60, bottom=60, left=80, right=80)
        p0 = c0.paragraphs[0]
        r_l = p0.add_run(f"{letter}) ")
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(11)
        r_l.font.bold = True
        r_s = p0.add_run(stmt_txt)
        r_s.font.name = "Times New Roman"
        r_s.font.size = Pt(11)
        
        c1 = tbl.cell(row_idx, 1)
        c1.width = Inches(0.9)
        c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c2 = tbl.cell(row_idx, 2)
        c2.width = Inches(0.9)
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def add_essay_question(doc, q_number, points_str, q_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.15)
    
    r_n = p.add_run(f"Câu {q_number} ({points_str}): ")
    r_n.font.name = "Times New Roman"
    r_n.font.size = Pt(11.5)
    r_n.font.bold = True
    
    r_t = p.add_run(q_text)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11.5)

def add_exam_footer_mark(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("---------- HẾT ----------\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r2 = p.add_run("Cán bộ coi thi không giải thích gì thêm. Thí sinh không được sử dụng tài liệu.")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10.5)
    r2.font.italic = True
    r2.font.color.rgb = GRAY_TEXT

def add_answers_section(doc, mcq_answers, essay_rubric=None, tf_answers=None):
    """
    Appends the answer key and grading rubric.
    """
    doc.add_page_break()
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    r_t = p_title.add_run("HƯỚNG DẪN CHẤM VÀ ĐÁP ÁN CHI TIẾT")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(14)
    r_t.font.bold = True
    r_t.font.color.rgb = NAVY_BLUE

    # MCQ Grid Table (e.g. 10 items per row)
    if mcq_answers:
        p_sec = doc.add_paragraph()
        r_sec = p_sec.add_run("I. ĐÁP ÁN TRẮC NGHIỆM")
        r_sec.font.name = "Times New Roman"
        r_sec.font.size = Pt(12)
        r_sec.font.bold = True
        
        # Determine number of columns (max 10 cols)
        items_per_row = 10 if len(mcq_answers) >= 10 else len(mcq_answers)
        rows_needed = (len(mcq_answers) + items_per_row - 1) // items_per_row
        
        tbl = doc.add_table(rows=rows_needed * 2, cols=items_per_row)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="718096", sz="4")
        
        for r_block in range(rows_needed):
            q_row = r_block * 2
            a_row = r_block * 2 + 1
            for c_idx in range(items_per_row):
                q_num = r_block * items_per_row + c_idx + 1
                if q_num <= len(mcq_answers):
                    # Question number row
                    cq = tbl.cell(q_row, c_idx)
                    cq.width = Inches(0.68)
                    set_cell_background(cq, "EDF2F7")
                    pq = cq.paragraphs[0]
                    pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    rq = pq.add_run(f"Câu {q_num}")
                    rq.font.name = "Times New Roman"
                    rq.font.size = Pt(10)
                    rq.font.bold = True
                    
                    # Answer row
                    ca = tbl.cell(a_row, c_idx)
                    ca.width = Inches(0.68)
                    pa = ca.paragraphs[0]
                    pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    ra = pa.add_run(mcq_answers[q_num - 1])
                    ra.font.name = "Times New Roman"
                    ra.font.size = Pt(11)
                    ra.font.bold = True
                    ra.font.color.rgb = CRIMSON

    # True/False answers
    if tf_answers:
        p_sec_tf = doc.add_paragraph()
        p_sec_tf.paragraph_format.space_before = Pt(8)
        r_sec_tf = p_sec_tf.add_run("II. ĐÁP ÁN TRẮC NGHIỆM ĐÚNG / SAI (DẠNG MỚI)")
        r_sec_tf.font.name = "Times New Roman"
        r_sec_tf.font.size = Pt(12)
        r_sec_tf.font.bold = True
        
        tbl_tf = doc.add_table(rows=len(tf_answers) + 1, cols=6)
        tbl_tf.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_tf, color="718096", sz="4")
        
        cols_cfg = [("Câu hỏi", Inches(1.2)), ("Lệnh a)", Inches(1.1)), ("Lệnh b)", Inches(1.1)), 
                    ("Lệnh c)", Inches(1.1)), ("Lệnh d)", Inches(1.1)), ("Điểm tối đa", Inches(1.4))]
        for c_idx, (col_name, w) in enumerate(cols_cfg):
            c = tbl_tf.cell(0, c_idx)
            c.width = w
            set_cell_background(c, "F1F5F9")
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(col_name)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.bold = True
            
        for row_idx, item in enumerate(tf_answers, start=1):
            q_lbl, a, b, c, d = item
            row_data = [q_lbl, a, b, c, d, "1.0 điểm"]
            for c_idx, val in enumerate(row_data):
                cell = tbl_tf.cell(row_idx, c_idx)
                cell.width = cols_cfg[c_idx][1]
                set_cell_padding(cell, top=60, bottom=60, left=60, right=60)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(val)
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)
                if val in ["Đúng", "Sai"]:
                    r.font.bold = True
                    r.font.color.rgb = CRIMSON if val == "Đúng" else NAVY_BLUE

        p_rule = doc.add_paragraph()
        p_rule.paragraph_format.space_before = Pt(4)
        r_rl = p_rule.add_run("* Quy tắc chấm điểm dạng trắc nghiệm Đúng/Sai của Bộ GD&ĐT:\n"
                              "- Thí sinh chỉ lựa chọn chính xác 01 ý: được 0,10 điểm.\n"
                              "- Thí sinh chỉ lựa chọn chính xác 02 ý: được 0,25 điểm.\n"
                              "- Thí sinh chỉ lựa chọn chính xác 03 ý: được 0,50 điểm.\n"
                              "- Thí sinh lựa chọn chính xác cả 04 ý: được 1,00 điểm.")
        r_rl.font.name = "Times New Roman"
        r_rl.font.size = Pt(10)
        r_rl.font.italic = True

    # Essay Rubric
    if essay_rubric:
        p_sec_es = doc.add_paragraph()
        p_sec_es.paragraph_format.space_before = Pt(8)
        r_sec_es = p_sec_es.add_run("III. HƯỚNG DẪN CHẤM TỰ LUẬN VÀ THANG ĐIỂM" if tf_answers else "II. HƯỚNG DẪN CHẤM TỰ LUẬN VÀ THANG ĐIỂM")
        r_sec_es.font.name = "Times New Roman"
        r_sec_es.font.size = Pt(12)
        r_sec_es.font.bold = True
        
        tbl_es = doc.add_table(rows=len(essay_rubric) + 1, cols=3)
        tbl_es.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_es, color="718096", sz="4")
        
        headers = [("Câu", Inches(1.0)), ("Nội dung đáp ứng yêu cầu cần đạt", Inches(5.0)), ("Điểm", Inches(1.0))]
        for c_idx, (col_name, w) in enumerate(headers):
            c = tbl_es.cell(0, c_idx)
            c.width = w
            set_cell_background(c, "F1F5F9")
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(col_name)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.bold = True
            
        for row_idx, (q_str, content_str, pt_str) in enumerate(essay_rubric, start=1):
            c0 = tbl_es.cell(row_idx, 0)
            c0.width = Inches(1.0)
            set_cell_padding(c0, top=60, bottom=60, left=60, right=60)
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(q_str)
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(10.5)
            r0.font.bold = True
            
            c1 = tbl_es.cell(row_idx, 1)
            c1.width = Inches(5.0)
            set_cell_padding(c1, top=60, bottom=60, left=80, right=80)
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(content_str)
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(10.5)
            
            c2 = tbl_es.cell(row_idx, 2)
            c2.width = Inches(1.0)
            set_cell_padding(c2, top=60, bottom=60, left=60, right=60)
            p2 = c2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(pt_str)
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(10.5)
            r2.font.bold = True

print("exam_builder_base loaded successfully.")
